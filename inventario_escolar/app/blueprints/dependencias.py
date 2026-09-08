import os
from flask import Blueprint, render_template, request, redirect, url_for, flash, current_app
from flask_login import login_required
from werkzeug.utils import secure_filename
from app.extensions import db
from app.models import Dependencia, Expediente
from app.decorators import roles_required, active_query

dependencias_bp = Blueprint('dependencias', __name__)

TIPOS = ['Practicas', 'Servicio', 'Ambos']

def obtener_sectores_dinamicos():
    sectores_bd = [s[0] for s in db.session.query(Dependencia.sector).filter(Dependencia.sector != None, db.func.trim(Dependencia.sector) != '').distinct().order_by(Dependencia.sector).all()]
    # Asegurar que siempre existan opciones base si se desea, aunque el usuario puede crear los suyos
    return sectores_bd

@dependencias_bp.before_request
@login_required
@roles_required('Super Admin', 'Practicas', 'Servicio')
def before_request():
    pass


@dependencias_bp.route('/')
def lista():
    tipo_filter = request.args.get('tipo_filter')
    search = request.args.get('search')

    query = active_query(Dependencia)
    if tipo_filter:
        query = query.filter_by(tipo=tipo_filter)
    if search:
        query = query.filter(Dependencia.nombre.ilike(f'%{search}%'))

    dependencias = query.order_by(Dependencia.nombre).all()
    return render_template('dependencias/lista.html',
                           dependencias=dependencias,
                           tipos=TIPOS,
                           sectores=obtener_sectores_dinamicos(),
                           tipo_filter=tipo_filter,
                           search=search)


@dependencias_bp.route('/crear', methods=['GET', 'POST'])
def crear():
    if request.method == 'POST':
        try:
            nombre = request.form.get('nombre').strip()
            tipo = request.form.get('tipo', 'Ambos')
            sector = request.form.get('sector')
            if sector:
                sector = sector.strip()
            domicilio = request.form.get('domicilio') or None
            contacto = request.form.get('contacto') or None
            telefono = request.form.get('telefono') or None
            correo = request.form.get('correo') or None

            if not nombre:
                flash('El nombre de la dependencia es obligatorio.', 'danger')
            else:
                dep = Dependencia(nombre=nombre, tipo=tipo, sector=sector,
                                  domicilio=domicilio, contacto=contacto,
                                  telefono=telefono, correo=correo)
                db.session.add(dep)
                db.session.commit()
                flash('Dependencia creada exitosamente.', 'success')
                return redirect(url_for('dependencias.lista'))
        except Exception as e:
            db.session.rollback()
            flash(f'Error: {str(e)}', 'danger')

    return render_template('dependencias/form.html',
                           dependencia=None,
                           tipos=TIPOS,
                           sectores=obtener_sectores_dinamicos())


@dependencias_bp.route('/<int:id>/editar', methods=['GET', 'POST'])
def editar(id):
    dependencia = active_query(Dependencia).filter_by(id=id).first_or_404()
    if request.method == 'POST':
        try:
            nombre = request.form.get('nombre').strip()
            if not nombre:
                flash('El nombre de la dependencia es obligatorio.', 'danger')
            else:
                dependencia.nombre = nombre
                dependencia.tipo = request.form.get('tipo', 'Ambos')
                sector = request.form.get('sector')
                dependencia.sector = sector.strip() if sector else None
                dependencia.domicilio = request.form.get('domicilio') or None
                dependencia.contacto = request.form.get('contacto') or None
                dependencia.telefono = request.form.get('telefono') or None
                dependencia.correo = request.form.get('correo') or None
                db.session.commit()
                flash('Dependencia actualizada.', 'success')
                return redirect(url_for('dependencias.lista'))
        except Exception as e:
            db.session.rollback()
            flash(f'Error: {str(e)}', 'danger')

    return render_template('dependencias/form.html',
                           dependencia=dependencia,
                           tipos=TIPOS,
                           sectores=obtener_sectores_dinamicos())


@dependencias_bp.route('/<int:id>/eliminar', methods=['POST'])
def eliminar(id):
    dependencia = active_query(Dependencia).filter_by(id=id).first_or_404()
    dependencia.is_deleted = True
    db.session.commit()
    flash('Dependencia eliminada.', 'success')
    return redirect(url_for('dependencias.lista'))


@dependencias_bp.route('/importar', methods=['GET', 'POST'])
def importar():
    resultado = None
    if request.method == 'POST':
        if 'archivo_excel' not in request.files:
            flash('No se subió ningún archivo.', 'danger')
            return redirect(request.url)

        file = request.files['archivo_excel']
        if file.filename == '':
            flash('No se seleccionó ningún archivo.', 'danger')
            return redirect(request.url)

        if file and file.filename.lower().endswith(('.xlsx', '.xls')):
            filepath = os.path.join('/tmp', secure_filename(file.filename))
            file.save(filepath)
            try:
                resultado = procesar_excel_dependencias(filepath)
            finally:
                if os.path.exists(filepath):
                    os.remove(filepath)
            if resultado['errores']:
                flash('Importación completada con algunos errores. Revisa el resumen.', 'warning')
            else:
                flash('Importación de dependencias completada.', 'success')
        else:
            flash('Formato de archivo no válido. Usa .xlsx o .xls', 'danger')

    return render_template(
        'practicas/importar.html',
        resultado=resultado,
        columnas_mapa=[],
        tipo_importacion='dependencias',
        modulo_label='Prácticas Profesionales',
        modulo_prefix='practicas',
    )


def procesar_excel_dependencias(filepath):
    """Importa la tabla de Dependencias de la hoja activa del libro."""
    import unicodedata
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filepath, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    except Exception as e:
        return {'insertados': 0, 'duplicados': 0, 'errores': [f'Error al leer el archivo: {str(e)}']}

    def normalizar(valor):
        texto = '' if valor is None else str(valor).strip().lower()
        texto = unicodedata.normalize('NFD', texto)
        texto = ''.join(c for c in texto if unicodedata.category(c) != 'Mn')
        return ''.join(c for c in texto if c.isalnum())

    encabezados = {
        'nombreempresa': 'nombre',
        'nombredelaempresa': 'nombre',
        'nombre': 'nombre',
        'empresa': 'nombre',
        'sector': 'sector',
        'ubicacion': 'domicilio',
        'domicilio': 'domicilio',
        'contacto': 'contacto',
        'telefono': 'telefono',
        'correo': 'correo',
        'email': 'correo',
    }
    header_row = None
    column_map = {}
    for row_idx, row in enumerate(rows[:30]):
        candidate = {}
        for col_idx, value in enumerate(row):
            field = encabezados.get(normalizar(value))
            if field and field not in candidate:
                candidate[field] = col_idx
        if 'nombre' in candidate:
            header_row = row_idx
            column_map = candidate
            break

    if header_row is None:
        return {'insertados': 0, 'duplicados': 0,
                'errores': ['No se encontró el encabezado NOMBRE DE LA EMPRESA en la hoja activa.']}

    def valor_fila(row, field):
        index = column_map.get(field)
        if index is None or index >= len(row) or row[index] is None:
            return None
        value = str(row[index]).strip()
        return value or None

    insertados = 0
    duplicados = 0
    errores = []
    tabla_iniciada = False

    for row_idx, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        fila_num = row_idx
        savepoint = db.session.begin_nested()
        try:
            nombre = valor_fila(row, 'nombre')
            if not nombre:
                if tabla_iniciada:
                    savepoint.rollback()
                    break
                savepoint.rollback()
                continue
            tabla_iniciada = True

            existente = Dependencia.query.filter(
                db.func.lower(Dependencia.nombre) == nombre.lower(),
                Dependencia.is_deleted == False,
            ).first()
            if existente:
                duplicados += 1
                savepoint.commit()
                continue

            sector = valor_fila(row, 'sector')
            domicilio = valor_fila(row, 'domicilio')
            contacto = valor_fila(row, 'contacto')
            telefono = valor_fila(row, 'telefono')
            correo = valor_fila(row, 'correo')

            dep = Dependencia(nombre=nombre, tipo='Practicas', sector=sector,
                              domicilio=domicilio, contacto=contacto,
                              telefono=telefono, correo=correo)
            db.session.add(dep)
            db.session.flush()
            insertados += 1
        except Exception as e:
            savepoint.rollback()
            errores.append(f'Fila {fila_num}: {str(e)}')

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        errores.append(f'Error al guardar: {str(e)}')

    return {'insertados': insertados, 'duplicados': duplicados, 'errores': errores}
